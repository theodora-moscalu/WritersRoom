import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from writersroom.importers.importer_factory import (
    ImporterFactory,
)
from writersroom.importers.pptx_importer import (
    PptxImporter,
)


def build_deck(path: Path):
    presentation = Presentation()

    title_layout = presentation.slide_layouts[0]
    bullet_layout = presentation.slide_layouts[1]

    cover = presentation.slides.add_slide(title_layout)
    cover.shapes.title.text = "THE WINE GAME"

    character = presentation.slides.add_slide(bullet_layout)
    character.shapes.title.text = "Nadia"
    body = character.placeholders[1].text_frame
    body.text = "Sommelier, 34"
    body.add_paragraph().text = "Wants respect from the old guard"
    character.notes_slide.notes_text_frame.text = (
        "Grew up above her father's failing wine shop."
    )

    table_slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )
    table_slide.shapes.title.text = "Season 1"
    table = table_slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(6), Inches(1)
    ).table
    table.cell(0, 0).text = "Question"
    table.cell(0, 1).text = "Can Nadia be trusted?"
    table.cell(1, 0).text = "Theme"
    table.cell(1, 1).text = "Belonging has a price."

    presentation.save(str(path))


def main():
    print("Testing PptxImporter...")

    path = Path(tempfile.mkdtemp()) / "bible.pptx"
    build_deck(path)

    importer = ImporterFactory.create(str(path))
    assert isinstance(importer, PptxImporter)

    result = importer.import_document(str(path))

    assert result.metadata["slides"] == "3"
    assert "THE WINE GAME" in result.text
    assert "TITLE: Nadia" in result.text
    assert "Sommelier, 34" in result.text
    assert "Wants respect from the old guard" in result.text
    assert "father's failing wine shop" in result.text
    assert "| Question | Can Nadia be trusted? |" in result.text

    print()
    print(result.text)
    print()
    print("PptxImporter tests passed.")


if __name__ == "__main__":
    main()
