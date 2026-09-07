from pathlib import Path

from pptx import Presentation

from writersroom.importers.base_importer import (
    BaseImporter,
)
from writersroom.importers.document_import_result import (
    DocumentImportResult,
)


class PptxImporter(BaseImporter):
    """Imports Microsoft PowerPoint decks (series bibles, character sheets)."""

    extensions = (
        ".pptx",
    )

    def import_document(
        self,
        path: str,
    ) -> DocumentImportResult:

        file = Path(path)

        presentation = Presentation(str(file))

        slides = []

        for index, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            slides.append(
                self._slide_text(index, slide)
            )

        return DocumentImportResult(
            filename=file.name,
            text="\n\n".join(slides),
            metadata={
                "format": "pptx",
                "slides": str(len(slides)),
            },
        )

    def _slide_text(self, index: int, slide) -> str:
        """Render one slide as structured flat text."""

        lines = [f"===== SLIDE {index} ====="]

        title = self._title(slide)

        if title:
            lines.append(f"TITLE: {title}")

        body = self._body(slide, title)

        if body:
            lines.append("BODY:")
            lines.extend(body)

        tables = self._tables(slide)

        if tables:
            lines.append("TABLE:")
            lines.extend(tables)

        notes = self._notes(slide)

        if notes:
            lines.append(f"NOTES: {notes}")

        return "\n".join(lines)

    def _title(self, slide) -> str:
        if (
            slide.shapes.title is not None
            and slide.shapes.title.has_text_frame
        ):
            return slide.shapes.title.text.strip()

        return ""

    def _body(self, slide, title: str) -> list[str]:
        lines = []

        for shape in slide.shapes:

            if (
                shape == slide.shapes.title
                or not shape.has_text_frame
            ):
                continue

            for paragraph in shape.text_frame.paragraphs:

                text = paragraph.text.strip()

                if not text or text == title:
                    continue

                indent = "  " * max(paragraph.level, 0)
                lines.append(f"{indent}- {text}")

        return lines

    def _tables(self, slide) -> list[str]:
        lines = []

        for shape in slide.shapes:

            if not shape.has_table:
                continue

            for row in shape.table.rows:

                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]

                lines.append("| " + " | ".join(cells) + " |")

        return lines

    def _notes(self, slide) -> str:
        if not slide.has_notes_slide:
            return ""

        frame = slide.notes_slide.notes_text_frame

        return (
            frame.text.strip().replace("\n", " ")
            if frame is not None
            else ""
        )
